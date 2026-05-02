from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
import chromadb
from sentence_transformers import SentenceTransformer
import os, hashlib, re
from typing import List, Dict, Any, Tuple

model = SentenceTransformer("intfloat/multilingual-e5-large")
client = chromadb.PersistentClient(path="./vectorstore")
collection = client.get_or_create_collection("laws")

def extract_article_num(text: str) -> Tuple[str,str]:
    article_match = re.search(r"Статья\s+([0-9]+(?:-[0-9]+)?)", text)
    article = article_match.group(1) if article_match else ""
    paragraph_match = re.search(r"(?:пункт|п\.)\s+([0-9]+)", text, re.IGNORECASE)
    paragraph = paragraph_match.group(1) if paragraph_match else ""
    return article, paragraph

def parse_docx(path: str) -> List[Dict[str, Any]]:
    doc = Document(path)
    fname = os.path.basename(path)
    all = []
    for para in doc.paragraphs:
        if para.text.strip():
            all.append(("paragraph", para.text.strip()))

    for table in doc.tables:
        table_text = parse_table(table)
        if table_text:
            all.append(("table", table_text))
    chunks = []
    cur_chunk= ""
    cur_article = ""
    cur_paragraph = ""
    chunk_cnt = 0
    for p_type, p_text in all:
        if p_text.startswith("Статья") and cur_chunk:
            chunks.append({
                "text": cur_chunk.strip(),
                "source": fname,
                "article": cur_article,
                "paragraph": cur_paragraph,
                "chunk_id": chunk_cnt
            })
            chunk_cnt +=1
            cur_chunk = ""
            article, paragraph = extract_article_num(p_text)
            cur_article = article if article else cur_article
            cur_paragraph = paragraph if paragraph else cur_paragraph
        else:
            article, paragraph = extract_article_num(p_text)
            if article:
                cur_article = article
            if paragraph:
                cur_paragraph = paragraph

        separator = "\n\n" if cur_chunk else ""
        new_chunk = cur_chunk + separator + p_text

        if len(new_chunk) > 1000 and cur_chunk:
            chunks.append({
                "text": cur_chunk.strip(),
                "source": fname,
                "article": cur_article,
                "paragraph": cur_paragraph,
                "chunk_id": chunk_cnt
            })
            chunk_cnt += 1
            overlap_text = cur_chunk[-200:] if len(cur_chunk) > 200 else cur_chunk
            cur_chunk = overlap_text + "\n\n" + p_text
        else:
            cur_chunk = new_chunk
    if cur_chunk:
        chunks.append({
            "text": cur_chunk.strip(),
            "source": fname,
            "article": cur_article,
            "paragraph": cur_paragraph,
            "chunk_id": chunk_cnt
        })

    return chunks

def parse_table(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        rows.append(" | ".join(cells))
    return "Таблица:\n" + "\n".join(rows)

def ingest_file(path):
    chunks = parse_docx(path)
    if not chunks:
        print(f"нет текста в файле")
        return
    
    ids = []
    docs = []
    metadatas = []
    texts_for_embedding = []
    for chunk in chunks:
        id = hashlib.md5(f"{chunk['source']}_{chunk['chunk_id']}".encode()).hexdigest()
        ids.append(id)
        docs.append(chunk["text"])
        metadatas.append({
            "source": chunk["source"],
            "article": chunk["article"],
            "paragraph": chunk["paragraph"],
            "chunk_id": chunk["chunk_id"]
        })
        texts_for_embedding.append(chunk["text"])
    embeddings = model.encode(texts_for_embedding, show_progress_bar=True).tolist()
    collection.upsert(
        documents=docs,
        embeddings=embeddings,
        ids=ids,
        metadatas=metadatas
    )
    print(f"загружено {len(chunks)} чанков из {path}")

def ingest_all(folder: str = "./data/laws"):
    if not os.path.exists(folder):
        print(f"папка {folder} не найдена, создаю...")
        os.makedirs(folder)
        return
    for f in os.listdir(folder):
        if f.lower().endswith(".docx"):
            ingest_file(os.path.join(folder, f))

if __name__ == "__main__":
    ingest_all()